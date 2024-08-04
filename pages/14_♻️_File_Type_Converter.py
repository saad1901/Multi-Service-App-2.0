import streamlit as st
from docx import Document
from fpdf import FPDF
import os

tab1, tab2, tab3 = st.tabs(["Docx to PDF converter", "test", "test2"])

class PDF(FPDF):
    def header(self):
        self.set_font('Arial', 'B', 12)
        self.cell(0, 10, 'Document Title', 0, 1, 'C')

    def footer(self):
        self.set_y(-15)
        self.set_font('Arial', 'I', 8)
        self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')

    def add_text(self, text, style):
        if 'B' in style:
            self.set_font('Arial', 'B', 12)
        elif 'I' in style:
            self.set_font('Arial', 'I', 12)
        elif 'U' in style:
            self.set_font('Arial', 'U', 12)
        else:
            self.set_font('Arial', '', 12)
        self.multi_cell(0, 10, text)
        self.ln()

def docx_to_pdf(docx_file_path, pdf_file_path):
    doc = Document(docx_file_path)
    pdf = PDF()
    pdf.add_page()

    for para in doc.paragraphs:
        runs = para.runs
        if runs:
            for run in runs:
                text = run.text
                style = ""
                if run.bold:
                    style += "B"
                if run.italic:
                    style += "I"
                if run.underline:
                    style += "U"
                pdf.add_text(text, style)
        else:
            pdf.ln()

    pdf.output(pdf_file_path)

with tab1:
    st.title("Batch DOCX to PDF Converter")

    uploaded_files = st.file_uploader("Choose DOCX files", type="docx", accept_multiple_files=True)

    if uploaded_files:
        with st.spinner('Converting...'):
            temp_dir = "temp_docs"
            os.makedirs(temp_dir, exist_ok=True)
            file_paths = []
            for uploaded_file in uploaded_files:
                file_path = os.path.join(temp_dir, uploaded_file.name)
                with open(file_path, "wb") as f:
                    f.write(uploaded_file.read())
                file_paths.append(file_path)

            pdf_files = []
            for file_path in file_paths:
                try:
                    pdf_file_path = os.path.splitext(file_path)[0] + '.pdf'
                    docx_to_pdf(file_path, pdf_file_path)
                    pdf_files.append(pdf_file_path)
                except Exception as e:
                    st.error(f"Error converting {file_path}: {str(e)}")

            st.success("Conversion successful!")
            for pdf_file in pdf_files:
                with open(pdf_file, "rb") as f:
                    st.download_button(
                        label=f"Download {os.path.basename(pdf_file)}",
                        data=f,
                        file_name=os.path.basename(pdf_file),
                        mime="application/pdf"
                    )
